"""Shared DOCX resume renderer.

Same profile dict as the PDF template so the two corpora stay comparable. DOCX
adds carriers a PDF does not have: hidden runs, comments, tracked changes and
image alt text, each of which is a placement.

Unicode also survives here, which the PDF path cannot manage with the standard
fonts: reportlab's Helvetica turns U+200B and Cyrillic look-alikes into "I". The
homoglyph and zero-width placements therefore live in this format.
"""

from __future__ import annotations

from dataclasses import dataclass
from pathlib import Path
from typing import Any

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor


@dataclass(frozen=True)
class DocxInjection:
    """Where the payload goes. Exactly one of these per attack."""

    placement: str
    payload: str


def _hide(run: Any) -> None:
    """w:vanish - the run is in the file but not in the reading view (ING-007)."""
    run._element.get_or_add_rPr().append(OxmlElement("w:vanish"))


def _add_comment_part(document: Any, text: str) -> None:
    """python-docx has no comments API, so the part is written directly."""
    from docx.opc.constants import RELATIONSHIP_TYPE as RT
    from docx.opc.packuri import PackURI
    from docx.opc.part import Part

    xml = (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        '<w:comments xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/'
        '2006/main">'
        '<w:comment w:id="1" w:author="Recruiting Operations" w:date="2026-03-04T10:00:00Z">'
        f"<w:p><w:r><w:t>{text}</w:t></w:r></w:p>"
        "</w:comment></w:comments>"
    )
    part = Part(
        PackURI("/word/comments.xml"),
        "application/vnd.openxmlformats-officedocument.wordprocessingml.comments+xml",
        xml.encode("utf-8"),
        document.part.package,
    )
    document.part.relate_to(part, RT.COMMENTS)


def _add_tracked_deletion(document: Any, text: str) -> None:
    """A deleted run: invisible to a reader, still present in the XML."""
    paragraph = document.add_paragraph()
    deletion = OxmlElement("w:del")
    deletion.set(qn("w:id"), "99")
    deletion.set(qn("w:author"), "Recruiting Operations")
    deletion.set(qn("w:date"), "2026-03-04T10:00:00Z")
    run = OxmlElement("w:r")
    del_text = OxmlElement("w:delText")
    del_text.text = text
    run.append(del_text)
    deletion.append(run)
    paragraph._p.append(deletion)


def _add_alt_text_image(document: Any, description: str) -> None:
    """Alt text on an inline shape: never rendered, always extractable."""
    import io

    from corpus.templates.png import solid_png

    document.add_picture(io.BytesIO(solid_png(60, 20, (240, 240, 240))))
    shape = document.inline_shapes[-1]
    shape._inline.docPr.set("descr", description)


# OPC core properties are capped at 255 characters each, and the payloads are
# longer than that. Splitting across fields keeps the whole payload in the file
# rather than truncating it: a half-payload would be a weaker attack than the
# one the manifest claims to be testing.
CORE_PROPERTY_LIMIT = 250
_CORE_CARRIERS = ("subject", "keywords", "comments", "category")


def _fill_core_properties(document: Any, payload: str) -> None:
    words = payload.split()
    chunks: list[str] = []
    current = ""
    for word in words:
        candidate = f"{current} {word}".strip()
        if len(candidate) > CORE_PROPERTY_LIMIT and current:
            chunks.append(current)
            current = word
        else:
            current = candidate
    if current:
        chunks.append(current)
    if len(chunks) > len(_CORE_CARRIERS):
        raise ValueError(
            f"payload needs {len(chunks)} core-property fields, "
            f"only {len(_CORE_CARRIERS)} available"
        )
    for field, chunk in zip(_CORE_CARRIERS, chunks, strict=False):
        setattr(document.core_properties, field, chunk)


def _body(document: Any, profile: dict[str, Any]) -> None:
    document.add_heading(profile["full_name"], level=0)
    contact = " | ".join(
        str(profile[key]) for key in ("headline", "location", "email", "phone")
        if profile.get(key)
    )
    document.add_paragraph(contact)
    if profile.get("portfolio_url"):
        document.add_paragraph(str(profile["portfolio_url"]))
    if profile.get("summary"):
        document.add_heading("Summary", level=1)
        document.add_paragraph(profile["summary"])
    if profile.get("skills"):
        document.add_heading("Skills", level=1)
        document.add_paragraph(", ".join(profile["skills"]))
    if profile.get("experience"):
        document.add_heading("Experience", level=1)
        for role in profile["experience"]:
            document.add_paragraph(
                f"{role['title']}, {role['company']} ({role['period']})"
            ).runs[0].bold = True
            for point in role.get("bullets", []):
                document.add_paragraph(point, style="List Bullet")
    if profile.get("education"):
        document.add_heading("Education", level=1)
        for item in profile["education"]:
            document.add_paragraph(
                f"{item['degree']} in {item['field']}, {item['institution']}"
                f" ({item.get('year', 'n/a')})"
            )


def render(
    profile: dict[str, Any],
    out_path: Path,
    *,
    injection: DocxInjection | None = None,
) -> Path:
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    document = docx.Document()
    document.core_properties.author = profile.get("full_name", "")
    document.core_properties.title = "Resume"

    _body(document, profile)

    if injection is not None:
        place, payload = injection.placement, injection.payload
        if place == "docx_hidden_run":
            run = document.add_paragraph().add_run(payload)
            _hide(run)
        elif place == "docx_white_text":
            run = document.add_paragraph().add_run(payload)
            run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        elif place == "docx_tiny_font":
            run = document.add_paragraph().add_run(payload)
            run.font.size = Pt(2)
        elif place == "docx_core_props":
            _fill_core_properties(document, payload)
        elif place == "docx_comment":
            _add_comment_part(document, payload)
        elif place == "docx_tracked_change":
            _add_tracked_deletion(document, payload)
        elif place == "docx_alt_text":
            _add_alt_text_image(document, payload)
        elif place == "docx_body":
            document.add_heading("Additional Information", level=1)
            document.add_paragraph(payload)
        else:
            raise ValueError(f"unknown docx placement {place!r}")

    document.save(str(out_path))
    return out_path
