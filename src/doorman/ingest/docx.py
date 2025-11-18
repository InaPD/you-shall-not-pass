"""DOCX parsing (spec 8.2).

DOCX carries several channels a PDF does not: comments, tracked changes and image
alt text all survive in the file while being invisible in a normal reading view,
and `w:vanish` hides a run outright. Each becomes a span or a metadata entry so
the ING-* rules and the classifier see them.

Runs are emitted with PDF-shaped span ids so `visible_text()` reconstructs reading
order the same way for both formats.
"""

from __future__ import annotations

import hashlib
from pathlib import Path
from typing import Any

import docx
from docx.oxml.ns import qn

from doorman.models import Document, Span

DEFAULT_FONT_PT = 11.0
CORE_FIELDS = (
    "title", "author", "subject", "keywords", "comments", "category", "last_modified_by",
)

_W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


# A style chain longer than this is a malformed document, not a deep hierarchy.
MAX_STYLE_DEPTH = 16


def _style_chain(style: Any) -> list[Any]:
    """A style and everything it is based on, outermost first.

    Word resolves a run's appearance through `w:basedOn` inheritance, so a style
    can set 2pt or near-white without the run carrying either property. Reading
    only direct formatting reports such a run as ordinary 11pt black text, which
    is a hidden-text construction ING-001 and ING-002 would never see.
    """
    chain: list[Any] = []
    seen: set[int] = set()
    while style is not None and id(style) not in seen and len(chain) < MAX_STYLE_DEPTH:
        seen.add(id(style))
        chain.append(style)
        style = getattr(style, "base_style", None)
    return chain


def _font_chain(run: Any, paragraph: Any) -> list[Any]:
    """Every font that could decide this run's appearance, nearest first:
    direct formatting, then the run's character style, then the paragraph's."""
    fonts = [run.font]
    for style in (getattr(run, "style", None), getattr(paragraph, "style", None)):
        fonts.extend(
            font for font in (getattr(s, "font", None) for s in _style_chain(style))
            if font is not None
        )
    return fonts


def _run_size(run: Any, paragraph: Any) -> float:
    for font in _font_chain(run, paragraph):
        size = getattr(font, "size", None)
        if size is not None:
            return float(size.pt)
    return DEFAULT_FONT_PT


def _run_color(run: Any, paragraph: Any) -> tuple[int, int, int]:
    """Falls back to black, which reads as visible.

    That is the conservative direction for a colour that cannot be resolved here
    - a theme colour carries no RGB - because the span is then treated as
    ordinary text rather than flagged. An unresolvable theme colour is a known
    blind spot for ING-001, recorded rather than papered over.
    """
    for font in _font_chain(run, paragraph):
        rgb = getattr(getattr(font, "color", None), "rgb", None)
        if rgb is None:
            continue
        value = str(rgb)
        try:
            return (int(value[0:2], 16), int(value[2:4], 16), int(value[4:6], 16))
        except ValueError:
            return (0, 0, 0)
    return (0, 0, 0)


def _is_hidden(run: Any) -> bool:
    """ING-007: w:vanish / w:specVanish marks a run as hidden text."""
    element = run._element
    properties = element.find(qn("w:rPr"))
    if properties is None:
        return False
    for tag in ("w:vanish", "w:specVanish"):
        node = properties.find(qn(tag))
        if node is not None and node.get(qn("w:val")) not in ("0", "false"):
            return True
    return False


def _spans(document: Any) -> list[Span]:
    spans: list[Span] = []
    for block_no, paragraph in enumerate(document.paragraphs):
        for run_no, run in enumerate(paragraph.runs):
            if not run.text:
                continue
            spans.append(
                Span(
                    id=f"p1-b{block_no}-l0-s{run_no}",
                    page=1,
                    text=run.text,
                    bbox=(0.0, 0.0, 0.0, 0.0),
                    size=_run_size(run, paragraph),
                    color_rgb=_run_color(run, paragraph),
                    # ING-007 is decided here because it is a property of the XML,
                    # not of geometry. hidden.py adds the rest.
                    hidden_reasons=["ING-007"] if _is_hidden(run) else [],
                )
            )
    return spans


def _part_xml(package: Any, name: str) -> bytes | None:
    for part in package.parts:
        if part.partname.endswith(name):
            return part.blob
    return None


def _comments(document: Any) -> dict[str, str]:
    blob = _part_xml(document.part.package, "comments.xml")
    if not blob:
        return {}
    from xml.etree import ElementTree

    try:
        root = ElementTree.fromstring(blob)
    except ElementTree.ParseError:
        return {}
    out: dict[str, str] = {}
    for index, comment in enumerate(root.findall(f"{_W}comment")):
        text = "".join(node.text or "" for node in comment.iter(f"{_W}t")).strip()
        if text:
            out[f"docx.comment.{index}"] = text
    return out


def _tracked_changes(document: Any) -> dict[str, str]:
    """Deleted runs survive in the XML: w:delText holds text a reader never sees."""
    out: dict[str, str] = {}
    body = document.element.body
    for kind, tag, text_tag in (("ins", f"{_W}ins", f"{_W}t"),
                                ("del", f"{_W}del", f"{_W}delText")):
        for index, node in enumerate(body.iter(tag)):
            text = "".join(child.text or "" for child in node.iter(text_tag)).strip()
            if text:
                out[f"docx.tracked.{kind}.{index}"] = text
    return out


def _alt_text(document: Any) -> dict[str, str]:
    out: dict[str, str] = {}
    for index, node in enumerate(document.element.body.iter()):
        if not node.tag.endswith("}docPr"):
            continue
        description = node.get("descr") or node.get("title") or ""
        if description.strip():
            out[f"docx.alt.{index}"] = description.strip()
    return out


def _core_properties(document: Any) -> dict[str, str]:
    properties = document.core_properties
    out: dict[str, str] = {}
    for field in CORE_FIELDS:
        value = getattr(properties, field, None)
        if value:
            out[f"docx.core.{field}"] = str(value)
    return out


def parse(path: Path) -> Document:
    """Parse a DOCX into a Document. Rule evaluation happens in `hidden.py`."""
    path = Path(path)
    raw = path.read_bytes()
    document = docx.Document(str(path))

    metadata: dict[str, str] = {}
    metadata.update(_core_properties(document))
    metadata.update(_comments(document))
    metadata.update(_tracked_changes(document))
    metadata.update(_alt_text(document))

    return Document(
        doc_id=path.stem,
        kind="docx",
        sha256=hashlib.sha256(raw).hexdigest(),
        page_count=1,  # DOCX has no fixed pagination until it is rendered
        spans=_spans(document),
        metadata=metadata,
    )
